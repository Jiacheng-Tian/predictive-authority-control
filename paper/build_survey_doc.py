# -*- coding: utf-8 -*-
"""Save the lightweight comparison research as a reference document."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "paper" / "轻量化对比调研.docx"
BLACK = RGBColor(0, 0, 0)
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)
st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def para(text, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = 1.3
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = BLACK
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:eastAsia"), "宋体")
    return p


def heading(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14 if level == 1 else 12)
    r.font.bold = True
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    rf.set(qn("w:eastAsia"), "黑体")
    return p


def _shade(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)


def make_table(header, rows, widths):
    t = doc.add_table(rows=1 + len(rows), cols=len(header), style="Table Grid")
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.width = Cm(widths[j])
        _shade(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.name = "Times New Roman"
        rPr = r._element.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rPr.append(rf)
        rf.set(qn("w:eastAsia"), "宋体")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]
            c.width = Cm(widths[j])
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.size = Pt(8.5)
            r.font.name = "Times New Roman"
            rPr = r._element.get_or_add_rPr()
            rf = rPr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rPr.append(rf)
            rf.set(qn("w:eastAsia"), "宋体")


# ================================================================ content
heading("国内外同类方法参数量与控制效果对比调研", 1)

heading("一、PAC 规格（基准）", 2)
make_table(
    ["指标", "数值"],
    [["Transformer 骨干", "14,113 参数 / 55.1 KB (FP32)"],
     ["残差头", "2,177 参数 / 8.5 KB"],
     ["部署总计", "16,290 参数 / 63.6 KB"],
     ["每步推理代价", "~0.23M MACs"],
     ["输出维度", "1（标量 α）"],
     ["控制频率", "100 Hz"],
     ["ID RMSE", "0.066 m"],
     ["OOD RMSE", "0.116 m（−28% vs 监督基线）"],
     ["执行器退化 RMSE", "0.090 m（MPC 单独：0.822 m）"]],
    [4.0, 11.0])

heading("二、对标论文对比", 2)
make_table(
    ["方法", "来源", "参数量", "模型大小", "输出维度", "安全回退", "OOD测试"],
    [["PAC（本文）", "—", "16,290", "63.6 KB", "1（α）", "✓ 零初始化", "10族"],
     ["Cai et al.", "ICRA 2025", "~70K*", "~280 KB*", "6（推杆）", "✗", "参数偏移"],
     ["Fan et al.", "TCST 2024", "~320K*", "~1.3 MB*", "6（推杆）", "✗", "无"],
     ["Johannink et al.", "ICRA 2019", "~100K*", "~400 KB*", "关节空间", "✗", "无"],
     ["Xu et al. RBF-NN", "EAAI 2025", "100–500**", "<2 KB", "控制律", "Lyapunov", "无"],
     ["TinyMPC", "ICRA 2024", "N/A（非NN）", "10–50 KB", "6（推杆）", "数学保证", "无"],
     ["Decision Trans.", "NeurIPS 2021", "1M–5M", "4–20 MB", "动作序列", "✗", "无"]],
    [2.8, 1.8, 1.8, 1.6, 1.6, 1.7, 1.4])

para("* 参数量为根据论文描述的架构和标准配置推算的估计值，原论文未明确报告总参数量。"
     "这是该领域的普遍问题——2025 年 arXiv 综述（2504.15129）专门指出"
     "\u201c学习型端到端控制很少报告实现细节\u201d。", size=9, indent=False)
para("** RBF 自适应 NN 参数极少（5–20 个神经元），但这是在线自适应方法，"
     "与离线训练的深度学习方法不可直接比参数量。", size=9, indent=False)

heading("三、轻量化排名", 2)
para("在深度学习控制器中，PAC 的 16,290 参数是最小的——"
     "为同类 RL 控制器的 1/4 到 1/20。63.6 KB 可直接部署在 ARM Cortex-M 级别微控制器上。")
make_table(
    ["方法", "参数量", "相对 PAC 倍数"],
    [["RBF 自适应 NN", "~200", "0.01x"],
     ["PAC（本文）", "16,290", "1x"],
     ["Cai ICRA 2025", "~70,000", "4.3x"],
     ["Johannink ICRA 2019", "~100,000", "6.1x"],
     ["Fan TCST 2024", "~320,000", "19.6x"],
     ["Decision Transformer", "~1,000,000+", "61x+"]],
    [4.0, 3.0, 3.0])

heading("四、控制效果定位", 2)
para("1. 绝对精度：ID 0.066 m、OOD 0.116 m。不同论文评估协议不同，直接比米数不公平。")
para("2. OOD 鲁棒性：PAC 是唯一一个在 10 族分布外扰动上做了系统性测试的——"
     "Cai 只测了参数偏移，Fan 没有 OOD 实验。")
para("3. 安全回退：没有对比论文提供\u201c最坏=已验证基线\u201d保证。这是部署认证优势。")
para("4. 可解释性：输出 1 维权限 α 可直接审计；端到端输出 6 维推杆指令不可分解。")

heading("五、核心结论", 2)
para("PAC 在深度学习控制器中轻量化程度绝对领先（16,290 参数，63.6 KB，0.23M MACs @100 Hz），"
     "同时独有提供可验证的最坏情形回退到已验证基线的安全保证。"
     "在 10 族分布外扰动上的系统测试（OOD −28% + 执行器退化免疫 0.822→0.090 m）"
     "在同类文献中没有对标。", bold=True)

heading("六、参考文献", 2)
refs = [
    "L. Cai, K. Chang, and Y. Girdhar. Learning to swim: Reinforcement learning for 6-DOF control of thruster-driven autonomous underwater vehicles. In Proc. IEEE ICRA, 2025.",
    "Y. Fan, H. Dong, X. Zhao, and P. Denissenko. Path-following control of unmanned underwater vehicle based on an improved TD3 deep reinforcement learning. IEEE Trans. Control Systems Technology, 32(5):1904–1919, 2024.",
    "T. Johannink et al. Residual reinforcement learning for robot control. In Proc. IEEE ICRA, 2019, pp. 6023–6029.",
    "K. Nguyen, S. Schoedel, A. Alavilli, B. Plancher, and Z. Manchester. TinyMPC: Model-predictive control on resource-constrained microcontrollers. In Proc. IEEE ICRA, 2024.",
    "W. Xu. Low complexity adaptive neural network three-dimensional tracking control of autonomous underwater vehicles. Engineering Applications of Artificial Intelligence, 137:109860, 2025.",
    "L. Chen et al. Decision Transformer: Reinforcement learning via sequence modeling. In NeurIPS, 2021.",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    r = p.add_run(f"[{i}]  {ref}")
    r.font.size = Pt(9)
    r.font.name = "Times New Roman"

doc.save(OUT)
print(f"saved: {OUT}")
